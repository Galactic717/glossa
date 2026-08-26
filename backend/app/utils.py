"""Document parsing, chunking and answer export."""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".txt", ".md", ".markdown", ".docx"}


class UnsupportedFormat(ValueError):
    pass


@dataclass
class Block:
    """A parsed region of a document, before chunking."""

    text: str
    page: int | None = None
    section: str | None = None


@dataclass
class Chunk:
    text: str
    index: int
    page: int | None = None
    section: str | None = None


def file_id(path: Path) -> str:
    """Content hash -- re-uploading the same file replaces it instead of duping."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()[:16]


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


# --------------------------------------------------------------------------- parsing


def _parse_pdf(path: Path) -> list[Block]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[Block] = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append(Block(text=text, page=number))
    if not blocks:
        raise UnsupportedFormat(
            "No extractable text in this PDF. It is probably a scan -- OCR it first."
        )
    return blocks


def _parse_docx(path: Path) -> list[Block]:
    import docx

    document = docx.Document(str(path))
    blocks: list[Block] = []
    section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            blocks.append(Block(text="\n".join(buffer), section=section))
            buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name.lower().startswith("heading"):
            flush()
            section = text
            continue
        buffer.append(text)
    flush()

    for table in document.tables:
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
            if any(cell.text.strip() for cell in row.cells)
        ]
        if rows:
            blocks.append(Block(text="\n".join(rows), section=section or "Table"))

    if not blocks:
        raise UnsupportedFormat("The .docx file contains no text.")
    return blocks


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")


def _parse_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        body = "\n".join(buffer).strip()
        if body:
            blocks.append(Block(text=body, section=section))
        buffer.clear()

    for line in text.splitlines():
        heading = _MD_HEADING.match(line)
        if heading:
            flush()
            section = heading.group(2).strip()
            continue
        buffer.append(line)
    flush()
    return blocks


def parse_file(path: Path) -> list[Block]:
    """Parse a document into blocks that carry page/section provenance."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in {".md", ".markdown"}:
        return _parse_markdown(path.read_text(encoding="utf-8", errors="replace"))
    if suffix == ".txt":
        text = path.read_text(encoding="utf-8", errors="replace").strip()
        if not text:
            raise UnsupportedFormat("The file is empty.")
        return [Block(text=text)]
    raise UnsupportedFormat(
        f"Unsupported format '{suffix}'. Supported: {', '.join(sorted(SUPPORTED_SUFFIXES))}."
    )


# --------------------------------------------------------------------------- chunking

_SPLIT = re.compile(r"(?<=[.!?…。？！])\s+|\n{2,}")


def _split_units(text: str, limit: int) -> list[str]:
    """Sentences/paragraphs, with over-long ones hard-wrapped at `limit`."""
    units: list[str] = []
    for piece in _SPLIT.split(text):
        piece = piece.strip()
        if not piece:
            continue
        while len(piece) > limit:
            units.append(piece[:limit])
            piece = piece[limit:]
        if piece:
            units.append(piece)
    return units


def chunk_blocks(blocks: list[Block], size: int, overlap: int) -> list[Chunk]:
    """Pack blocks into ~`size`-char chunks with `overlap` chars of tail carry-over.

    Chunks never span two blocks, so a citation always points at one real
    page or heading rather than a seam between two of them.
    """
    overlap = max(0, min(overlap, size // 2))
    chunks: list[Chunk] = []
    index = 0

    for block in blocks:
        current = ""
        for unit in _split_units(block.text, size):
            candidate = f"{current} {unit}".strip() if current else unit
            if len(candidate) <= size:
                current = candidate
                continue
            if current:
                chunks.append(
                    Chunk(text=current, index=index, page=block.page, section=block.section)
                )
                index += 1
                current = f"{current[-overlap:]} {unit}".strip() if overlap else unit
            else:
                current = unit
        if current.strip():
            chunks.append(
                Chunk(text=current.strip(), index=index, page=block.page, section=block.section)
            )
            index += 1

    return chunks


# --------------------------------------------------------------------------- export


def answer_to_markdown(answer: dict) -> str:
    lines = [
        f"# {answer['question']}",
        "",
        answer["answer"],
        "",
        "## Джерела / Sources",
        "",
    ]
    for source in answer.get("sources", []):
        where = source.get("location") or ""
        lines.append(f"**[{source['ref']}] {source['filename']}**{f' — {where}' if where else ''}")
        lines.append("")
        lines.append("> " + source["text"].replace("\n", "\n> "))
        lines.append("")
    meta = answer.get("meta", {})
    if meta:
        lines.append("---")
        lines.append(
            "*"
            + " · ".join(f"{key}: {value}" for key, value in meta.items())
            + "*"
        )
    return "\n".join(lines)


def answer_to_json(answer: dict) -> str:
    return json.dumps(answer, ensure_ascii=False, indent=2)
